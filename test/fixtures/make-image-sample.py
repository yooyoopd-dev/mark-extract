#!/usr/bin/env python3
"""그림이 든 시험 자료 생성 (build.30).

`sample-ko.*` · `sample-table.*` 에는 **그림이 한 장도 없었다.** 그래서 네 어댑터가
모두 죽은 파일 참조를 내보내는 동안 단언 60여 개가 전부 통과했다 — build.8 의 표,
build.25 의 별표와 같은 종류의 구멍이다.

    pip install pillow python-docx python-pptx
    python3 test/fixtures/make-image-sample.py

PDF 는 만들어진 HTML 을 헤드리스 크로미움으로 찍는다.

    chrome --headless --disable-gpu --no-pdf-header-footer \
           --print-to-pdf=sample-image.pdf file://$PWD/sample-image.html

그림을 **두 쪽에 걸쳐** 두는 것이 요점이다. 한 쪽에만 있으면 쪽 번호가 늘 1 이라
위치 표시가 맞는지 틀리는지 구별되지 않는다. 표 셀 안에도 한 장 둔다 — 그 자리에는
인용 줄을 넣을 수 없어 처리 경로가 다르다.
"""
import base64
import io
from pathlib import Path

from PIL import Image, ImageDraw

HERE = Path(__file__).parent
TITLE = "그림이 든 시험 자료"


def png(label: str, color: tuple[int, int, int]) -> bytes:
    """글자가 든 단색 사각형. 내용이 아니라 '그림이 거기 있었다'는 사실만 필요하다."""
    size = (320, 180)
    im = Image.new("RGB", size, color)
    draw = ImageDraw.Draw(im)
    draw.rectangle([4, 4, size[0] - 5, size[1] - 5], outline=(20, 20, 20), width=3)
    draw.text((20, 80), label, fill=(10, 10, 10))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


FIG1 = png("FIGURE ONE", (210, 226, 246))
FIG2 = png("FIGURE TWO", (246, 226, 210))
FIG3 = png("FIGURE THREE", (218, 240, 218))


def write_html() -> None:
    b64 = [base64.b64encode(raw).decode() for raw in (FIG1, FIG3, FIG2)]
    (HERE / "sample-image.html").write_text(
        f"""<!doctype html>
<html lang="ko">
<meta charset="utf-8">
<title>{TITLE}</title>
<style>
  @page {{ size: A4; margin: 20mm; }}
  body {{ font-family: "Noto Sans CJK KR", "Noto Sans KR", sans-serif; font-size: 12pt; line-height: 1.6; }}
  h1 {{ font-size: 20pt; }}
  h2 {{ font-size: 15pt; margin-top: 1.4em; }}
  img {{ display: block; margin: 12px 0; width: 80mm; }}
  table {{ border-collapse: collapse; }}
  td, th {{ border: 1px solid #333; padding: 6px 10px; }}
  .page-break {{ break-before: page; }}
</style>

<h1>{TITLE}</h1>

<p>이 문서는 못 읽은 그림의 위치 표시를 검증하려고 만든 자료다. 그림을 두 쪽에 걸쳐
배치해 쪽 번호가 실제로 갈리는지 본다.</p>

<h2>1. 첫째 쪽 그림</h2>

<p>아래 그림은 첫째 쪽에 있다.</p>

<img src="data:image/png;base64,{b64[0]}" alt="첫째 그림">

<p>그림 다음에 오는 문단이다. 위치 표시 줄이 이 문단을 삼키지 않아야 한다.</p>

<h2>2. 표 안의 그림</h2>

<table>
  <tr><th>구분</th><th>그림</th></tr>
  <tr><td>표 안</td><td><img src="data:image/png;base64,{b64[1]}" alt="표 안 그림"></td></tr>
</table>

<h2 class="page-break">3. 둘째 쪽 그림</h2>

<p>아래 그림은 둘째 쪽에 있다. 쪽 번호가 1 이 아니라 2 로 나와야 한다.</p>

<img src="data:image/png;base64,{b64[2]}" alt="둘째 그림">

<p>문서 끝 문단이다.</p>
</html>
""",
        encoding="utf-8",
    )


def write_docx() -> None:
    from docx import Document
    from docx.shared import Mm

    doc = Document()
    doc.add_heading(TITLE, level=1)
    doc.add_paragraph("이 문서는 못 읽은 그림의 위치 표시를 검증하려고 만든 자료다.")

    doc.add_heading("1. 첫째 그림", level=2)
    doc.add_paragraph("아래 그림은 첫 번째 절에 있다.")
    doc.add_picture(io.BytesIO(FIG1), width=Mm(80))
    doc.add_paragraph("그림 다음에 오는 문단이다.")

    doc.add_heading("2. 표 안의 그림", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "구분"
    table.cell(0, 1).text = "그림"
    table.cell(1, 0).text = "표 안"
    table.cell(1, 1).paragraphs[0].add_run().add_picture(io.BytesIO(FIG3), width=Mm(50))

    doc.add_heading("3. 둘째 그림", level=2)
    doc.add_picture(io.BytesIO(FIG2), width=Mm(80))
    doc.save(HERE / "sample-image.docx")


def write_pptx() -> None:
    from pptx import Presentation
    from pptx.util import Mm

    deck = Presentation()
    blank = deck.slide_layouts[6]
    title_only = deck.slide_layouts[5]

    first = deck.slides.add_slide(title_only)
    first.shapes.title.text = "1. 첫째 슬라이드"
    first.shapes.add_picture(io.BytesIO(FIG1), Mm(20), Mm(60), width=Mm(80))

    second = deck.slides.add_slide(blank)
    second.shapes.add_picture(io.BytesIO(FIG2), Mm(20), Mm(40), width=Mm(80))

    deck.save(HERE / "sample-image.pptx")


if __name__ == "__main__":
    write_html()
    write_docx()
    write_pptx()
    print("만들었습니다: sample-image.html · sample-image.docx · sample-image.pptx")
    print("PDF 는 README 의 크로미움 명령으로 sample-image.html 을 찍으세요.")
