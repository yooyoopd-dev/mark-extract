#!/usr/bin/env python3
"""시험 자료(DOCX·XLSX·XLS·PPTX) 생성.

PDF 와 마찬가지로 산출물을 커밋해 두되(CI 가 플랫폼마다 생성기를 갖출 수 없다)
만드는 방법을 함께 남겨 재현할 수 있게 한다.

    pip install python-docx openpyxl python-pptx xlwt
    python3 test/fixtures/make-office.py

내용은 sample-ko.pdf 와 같은 뼈대다 — 한글·영문 혼재, 제목 계층, 표, 목록,
숫자·특수문자. 포맷이 달라도 같은 것을 확인할 수 있어야 비교가 된다.
"""
from pathlib import Path

HERE = Path(__file__).parent
TITLE = "문서 변환 시험 자료"
INTRO = ("이 문서는 Mark Extract의 어댑터를 시험하기 위해 만든 자료다. "
         "한국어와 English가 섞여 있으며, 제목 계층과 표가 마크다운으로 어떻게 넘어오는지 확인한다.")
TABLE = [
    ["형식", "엔진", "비고"],
    ["PDF", "opendataloader-pdf", "레이아웃 분석"],
    ["DOCX", "kordoc", "순수 JS 경로"],
    ["PPTX", "자체 구현", "markitdown 규칙"],
]
NUMBERS = "금액 1,234,567원 · 비율 45.6% · 날짜 2026-09-16 · 괄호 (참고) · 따옴표 “인용”"


def docx_file():
    from docx import Document
    d = Document()
    d.add_heading(TITLE, level=1)
    d.add_paragraph(INTRO)
    d.add_heading("1. 제목 계층", level=2)
    d.add_paragraph("본문 문단이다. 레이아웃 분석이 제대로 되면 이 문단은 위 제목에 속한 것으로 읽힌다.")
    d.add_heading("1.1 하위 절", level=3)
    d.add_paragraph("세 번째 수준의 제목 아래 문단. Mixed content with English words like conversion, adapter, and markdown.")
    d.add_heading("2. 표", level=2)
    t = d.add_table(rows=0, cols=3)
    t.style = "Table Grid"
    for row in TABLE:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    d.add_heading("3. 목록", level=2)
    for item in ["첫 번째 항목", "두 번째 항목 — 특수문자 · 포함", "세 번째 항목"]:
        d.add_paragraph(item, style="List Bullet")
    d.add_heading("4. 숫자와 기호", level=2)
    d.add_paragraph(NUMBERS)
    d.save(HERE / "sample-ko.docx")


def xlsx_file():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "변환 대상"
    ws["A1"] = TITLE
    for r, row in enumerate(TABLE, start=3):
        for c, text in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=text)
    ws["A8"] = NUMBERS
    second = wb.create_sheet("두 번째 시트")
    second["A1"] = "시트가 둘 이상일 때 모두 나오는지 본다."
    second["A2"] = "English and 한국어 mixed."
    wb.save(HERE / "sample-ko.xlsx")


def xls_file():
    import xlwt
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("변환 대상")
    ws.write(0, 0, TITLE)
    for r, row in enumerate(TABLE, start=2):
        for c, text in enumerate(row):
            ws.write(r, c, text)
    ws.write(7, 0, NUMBERS)
    wb.save(str(HERE / "sample-ko.xls"))


def pptx_file():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    p = Presentation()

    s1 = p.slides.add_slide(p.slide_layouts[1])
    s1.shapes.title.text = TITLE
    s1.placeholders[1].text = INTRO
    s1.notes_slide.notes_text_frame.text = "발표자 노트다. 이 줄이 마크다운에 나와야 한다."

    s2 = p.slides.add_slide(p.slide_layouts[5])
    s2.shapes.title.text = "2. 표"
    shape = s2.shapes.add_table(len(TABLE), 3, Inches(0.8), Inches(1.8), Inches(8), Inches(2))
    for r, row in enumerate(TABLE):
        for c, text in enumerate(row):
            shape.table.cell(r, c).text = text

    s3 = p.slides.add_slide(p.slide_layouts[5])
    s3.shapes.title.text = "3. 목록과 숫자"
    box = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8), Inches(3))
    tf = box.text_frame
    tf.text = "첫 번째 항목"
    for text in ["두 번째 항목 — 특수문자 · 포함", "세 번째 항목", NUMBERS]:
        tf.add_paragraph().text = text
    for para in tf.paragraphs:
        para.font.size = Pt(18)

    p.save(HERE / "sample-ko.pptx")




# ── 복잡한 표 ────────────────────────────────────────────────────────────
#
# build.8 Windows 실측에서 드러난 결함을 재현하는 자료다. sample-ko.* 의 표는
# 3×3 단순 표라 파이프 표로 잘 넘어오고, 그래서 단언 52개가 다 통과하는데도
# 실제 문서에서는 깨졌다.
#
# 여기서 노리는 것:
#   - 병합 셀 → kordoc 이 tableToMarkdown() 대신 tableToHtml() 로 빠진다
#   - 셀 안 여러 줄 → 셀 내용이 <br> 로 이어진다
#   - 셀 안 | < > & → 표 구조를 깨거나 엔티티로 이스케이프된다

TABLE_TITLE = "병합 표 시험 자료"


def table_docx_file():
    from docx import Document

    d = Document()
    d.add_heading(TABLE_TITLE, level=1)
    d.add_paragraph("병합 셀과 셀 안 줄바꿈이 마크다운으로 어떻게 넘어오는지 본다.")

    d.add_heading("1. 가로·세로 병합", level=2)
    # 4행 3열. (0,0)-(0,2) 가로 병합, (1,0)-(2,0) 세로 병합.
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"

    head = t.rows[0].cells
    head[0].merge(head[2]).text = "2026년 추진 계획"

    t.cell(1, 0).merge(t.cell(2, 0)).text = "1분기"
    t.cell(1, 1).text = "목표일정"
    t.cell(1, 2).text = "비고"
    t.cell(2, 1).text = "설계 확정"
    t.cell(2, 2).text = "완료"

    t.cell(3, 0).text = "2분기"
    t.cell(3, 1).text = "구현"
    t.cell(3, 2).text = "진행"

    d.add_heading("2. 셀 안 여러 줄과 특수문자", level=2)
    t2 = d.add_table(rows=3, cols=2)
    t2.style = "Table Grid"
    t2.cell(0, 0).text = "항목"
    t2.cell(0, 1).text = "내용"

    # 한 셀 안의 여러 문단 → 마크다운에서 <br> 가 된다
    cell = t2.cell(1, 0)
    cell.text = "첫째 줄"
    cell.add_paragraph("둘째 줄")
    cell.add_paragraph("셋째 줄")
    # 문구를 고를 때 kordoc 의 자간 복원 휴리스틱을 피한다 — 30자 이하이고 한글
    # 1글자 토큰이 70% 이상이면 공백을 전부 지운다 (README 참조). "셀 하나에 세 줄"
    # 은 4토큰 중 3개가 단음절이라 "셀하나에세줄" 이 된다.
    t2.cell(1, 1).text = "이 셀에는 문단이 세 개 들어 있다"

    # 표 문법과 HTML 을 깨뜨릴 수 있는 글자들
    t2.cell(2, 0).text = "특수문자"
    t2.cell(2, 1).text = "파이프 | 와 꺾쇠 <태그> 와 앰퍼샌드 & 를 담았다"

    d.add_paragraph("표 밖 문단이다. 여기에는 줄바꿈 표시가 나오면 안 된다.")
    # 표 밖 특수문자. 엔진이 이것을 &lt; &gt; &amp; 로 이스케이프하므로 표 변환과는
    # 별개인 디코드 경로를 지나간다. 표 안에만 두면 xmldom 이 파싱하며 풀어 버려서
    # 디코드가 빠져도 검증이 통과한다.
    d.add_paragraph("표 밖 특수문자: 꺾쇠 <태그> · 앰퍼샌드 & · 부등호 5 < 10 이다.")
    d.save(HERE / "sample-table.docx")


for make in (docx_file, xlsx_file, xls_file, pptx_file, table_docx_file):
    make()
    print(f"생성: {make.__name__}")
